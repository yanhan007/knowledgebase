"""文档加载器模块

支持格式：PDF、DOCX、TXT、MD
特性：批量处理、错误处理、处理结果统计
"""

import logging
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Set

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


@dataclass
class LoadResult:
    """批量加载结果统计"""

    success_count: int = 0
    fail_count: int = 0
    total_chunks: int = 0
    errors: List[str] = field(default_factory=list)

    def add_success(self, chunks: int):
        self.success_count += 1
        self.total_chunks += chunks

    def add_fail(self, file_path: str, error: str):
        self.fail_count += 1
        self.errors.append(f"{file_path}: {error}")

    @property
    def summary(self) -> str:
        return (
            f"加载完成: 成功 {self.success_count} 个, "
            f"失败 {self.fail_count} 个, "
            f"共 {self.total_chunks} 个文本块"
        )


class BaseParser(ABC):
    """文档解析器基类"""

    @abstractmethod
    def parse(self, file_path: Path) -> List[Document]:
        """解析文档，返回文档块列表"""
        pass


class TextParser(BaseParser):
    """TXT/MD 文本文件解析器"""

    SUPPORTED_ENCODINGS = ["utf-8", "gbk", "gb2312", "latin-1"]

    def parse(self, file_path: Path) -> List[Document]:
        content = self._read_with_fallback_encoding(file_path)
        return [
            Document(
                page_content=content,
                metadata={"source": str(file_path), "type": file_path.suffix},
            )
        ]

    def _read_with_fallback_encoding(self, file_path: Path) -> str:
        """尝试多种编码读取文件"""
        for encoding in self.SUPPORTED_ENCODINGS:
            try:
                return file_path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"无法解码文件，尝试过的编码: {self.SUPPORTED_ENCODINGS}")


class PDFParser(BaseParser):
    """PDF 文件解析器"""

    def parse(self, file_path: Path) -> List[Document]:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise ImportError("请安装 PyPDF2: pip install PyPDF2")

        reader = PdfReader(str(file_path))
        documents = []

        for page_num, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
                if text and text.strip():
                    documents.append(
                        Document(
                            page_content=text.strip(),
                            metadata={
                                "source": str(file_path),
                                "page": page_num,
                                "type": "pdf",
                            },
                        )
                    )
            except Exception as e:
                logger.warning(f"PDF 第 {page_num + 1} 页解析失败: {e}")

        if not documents:
            raise ValueError("PDF 文件无法提取任何文本内容")

        return documents


class DocxParser(BaseParser):
    """Word 文档解析器"""

    def parse(self, file_path: Path) -> List[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

        doc = DocxDocument(str(file_path))
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        if not paragraphs:
            raise ValueError("Word 文档无有效文本内容")

        content = "\n\n".join(paragraphs)
        return [
            Document(
                page_content=content,
                metadata={"source": str(file_path), "type": "docx"},
            )
        ]


class DocumentLoader:
    """文档加载器主类"""

    DEFAULT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

    def __init__(self, extensions: Optional[Set[str]] = None):
        self.extensions = extensions or self.DEFAULT_EXTENSIONS
        self._parsers = {
            ".txt": TextParser(),
            ".md": TextParser(),
            ".pdf": PDFParser(),
            ".docx": DocxParser(),
        }

    def load(self, file_path: str) -> List[Document]:
        """加载单个文档"""
        path = self._validate_file(file_path)
        parser = self._get_parser(path)
        return parser.parse(path)

    def load_batch(self, file_paths: List[str]) -> tuple[List[Document], LoadResult]:
        """批量加载文档，返回 (文档列表, 加载结果)"""
        result = LoadResult()
        all_documents = []

        for file_path in file_paths:
            try:
                docs = self.load(file_path)
                all_documents.extend(docs)
                result.add_success(len(docs))
                logger.info(f"成功加载: {file_path} ({len(docs)} 块)")
            except Exception as e:
                result.add_fail(file_path, str(e))
                logger.error(f"加载失败: {file_path} - {e}")

        logger.info(result.summary)
        return all_documents, result

    def load_directory(self, dir_path: str) -> tuple[List[Document], LoadResult]:
        """加载目录下所有支持的文档"""
        path = Path(dir_path)
        if not path.is_dir():
            raise NotADirectoryError(f"目录不存在: {dir_path}")

        file_paths = []
        for ext in self.extensions:
            file_paths.extend(str(p) for p in path.rglob(f"*{ext}"))
            file_paths.extend(str(p) for p in path.rglob(f"*{ext.upper()}"))

        # 去重
        file_paths = list(set(file_paths))
        logger.info(f"发现 {len(file_paths)} 个文档")

        return self.load_batch(file_paths)

    def _validate_file(self, file_path: str) -> Path:
        """验证文件"""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        if not path.is_file():
            raise ValueError(f"不是文件: {file_path}")
        if path.suffix.lower() not in self.extensions:
            raise ValueError(f"不支持的文件格式: {path.suffix}")
        return path

    def _get_parser(self, path: Path) -> BaseParser:
        """获取对应的解析器"""
        ext = path.suffix.lower()
        parser = self._parsers.get(ext)
        if not parser:
            raise ValueError(f"未找到 {ext} 格式的解析器")
        return parser
